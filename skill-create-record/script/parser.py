from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Any

import openpyxl


@dataclass
class RawFormData:
    page1_text: str
    page2_text: str
    page3_questions: list[str] = field(default_factory=list)


def excel_date_to_iso(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%dT00:00:00Z")
    if isinstance(value, (int, float)):
        base = datetime(1899, 12, 30)
        dt = base + timedelta(days=int(value))
        return dt.strftime("%Y-%m-%dT00:00:00Z")
    return str(value)


def _row_to_line(row) -> str:
    is_birthday_row = any(
        hasattr(c, "value") and isinstance(c.value, str) and "出生年月日" in c.value
        for c in row
    )

    cells = []
    for c in row:
        if hasattr(c, "value"):
            v = c.value
            if v is None:
                continue
            if isinstance(v, datetime):
                cells.append(excel_date_to_iso(v) or "")
            elif isinstance(v, (int, float)) and getattr(c, "is_date", False):
                cells.append(excel_date_to_iso(v) or "")
            elif is_birthday_row and isinstance(v, (int, float)) and not isinstance(v, bool):
                cells.append(excel_date_to_iso(v) or str(v))
            else:
                text = str(v).strip()
                if text:
                    cells.append(text)
        else:
            if c is None:
                continue
            if isinstance(c, datetime):
                cells.append(excel_date_to_iso(c) or "")
            else:
                text = str(c).strip()
                if text:
                    cells.append(text)
    return " | ".join(cells)


def parse_xlsx(path: str) -> RawFormData:
    wb = openpyxl.load_workbook(path, data_only=True)

    ws1 = wb["Page１"]
    page1_lines: list[str] = []
    for row in ws1.iter_rows():
        line = _row_to_line(row)
        if line:
            page1_lines.append(line)

    ws2 = wb["Page2"]
    page2_lines: list[str] = []
    for row in ws2.iter_rows(values_only=True):
        parts = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if parts:
            page2_lines.append(" ".join(parts))

    ws3 = wb["Page3"]
    questions: list[str] = []
    for row in ws3.iter_rows(values_only=True):
        first = row[0] if row else None
        if isinstance(first, int) and first >= 1:
            text = " ".join(str(c) for c in row[1:] if c is not None and str(c).strip())
            if text.strip():
                questions.append(text.strip())

    return RawFormData(
        page1_text="\n".join(page1_lines),
        page2_text="\n".join(page2_lines),
        page3_questions=questions,
    )
